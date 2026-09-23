"""Genera la plantilla de nómina para el 2276 a partir de config/plantilla_nomina.yaml.

    python ejemplos/generar_plantilla_nomina.py   -> ejemplos/plantilla_nomina_2276.xlsx
"""
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

AQUI = Path(__file__).resolve().parent
COLUMNAS = yaml.safe_load((AQUI.parent / "config" / "plantilla_nomina.yaml").read_text(encoding="utf-8"))["columnas"]
TEXTO = {"nit", "tipo documento", "primer apellido", "segundo apellido", "primer nombre", "otros nombres", "direccion",
         "ciudad", "departamento", "pais", "tipo_doc_dependiente", "nit_dependiente", "id_fideicomiso",
         "tipo_doc_colaboracion", "nit_colaboracion"}
FILAS = 500

NOTAS = [
    "Una fila por empleado o beneficiario con los ACUMULADOS del año gravable (enero a diciembre).",
    "Tome los valores del reporte de nómina electrónica o de los acumulados por concepto del software de nómina.",
    "No cambie ni reordene los encabezados de la hoja Nomina: el aplicativo los lee por nombre.",
    "Valores en pesos, sin decimales, sin puntos de miles ni signo $. Deje vacío lo que no aplique.",
    "No agregue filas de totales ni de subtotales.",
    "Los honorarios y servicios de independientes solo van aquí si se les aplicó la tabla del art. 383; si no, van en el 1001.",
    "El aplicativo compara los salarios con las cuentas contables de sueldos y deja una ALERTA si no cuadran.",
    "Las ayudas marcadas 'Práctica' son criterio profesional: confírmelas con el contador responsable.",
]


def main() -> Path:
    wb = Workbook()
    ws = wb.active
    ws.title = "Nomina"
    azul, gris = PatternFill("solid", fgColor="1F4E78"), PatternFill("solid", fgColor="DDEBF7")
    for j, c in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=j, value=c["columna"])
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = azul if c.get("obligatorio") else PatternFill("solid", fgColor="2E75B6")
        celda.alignment = Alignment(wrap_text=True, vertical="center")
        nota = f"2276: {c['dian']}" + (f"\n\n{c['ayuda']}" if c.get("ayuda") else "")
        celda.comment = Comment(nota, "Exógena YC", width=320, height=160)
        letra = celda.column_letter
        ws.column_dimensions[letra].width = 16 if c["columna"] not in TEXTO else 18
        if c["columna"] not in TEXTO:
            for i in range(2, FILAS + 2):
                ws.cell(row=i, column=j).number_format = "#,##0"
        else:
            for i in range(2, FILAS + 2):
                ws.cell(row=i, column=j).number_format = "@"
    ws.row_dimensions[1].height = 32
    ws.freeze_panes = "B2"
    tipos = DataValidation(type="list", formula1='"CC,CE,TI,PPT,PEP,PA,NIT,RC"', allow_blank=True)
    tipos.error, tipos.errorTitle = "Use CC, CE, TI, PPT, PEP, PA, NIT o RC", "Tipo de documento"
    ws.add_data_validation(tipos)
    tipos.add(f"B2:B{FILAS + 1}")

    ins = wb.create_sheet("Instrucciones")
    ins.column_dimensions["A"].width = 30
    ins.column_dimensions["B"].width = 60
    ins.column_dimensions["C"].width = 90
    ins.column_dimensions["D"].width = 12
    ins["A1"] = "Plantilla de nómina para el formato 2276 v4 (rentas de trabajo y pensiones)"
    ins["A1"].font = Font(bold=True, size=13)
    fila = 3
    for n in NOTAS:
        ins.cell(row=fila, column=1, value="• " + n)
        fila += 1
    fila += 1
    for j, t in enumerate(["Columna de la plantilla", "Columna del prevalidador 2276", "Qué va", "Obligatoria"], start=1):
        c = ins.cell(row=fila, column=j, value=t)
        c.font, c.fill = Font(bold=True), gris
    for c in COLUMNAS:
        fila += 1
        ins.cell(row=fila, column=1, value=c["columna"])
        ins.cell(row=fila, column=2, value=c["dian"]).alignment = Alignment(wrap_text=True, vertical="top")
        ins.cell(row=fila, column=3, value=c.get("ayuda", "")).alignment = Alignment(wrap_text=True, vertical="top")
        ins.cell(row=fila, column=4, value="Sí" if c.get("obligatorio") else "")

    salida = AQUI / "plantilla_nomina_2276.xlsx"
    wb.save(salida)
    return salida


if __name__ == "__main__":
    print(main())
