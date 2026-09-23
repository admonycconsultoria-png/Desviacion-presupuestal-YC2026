"""Escritura de formatos (orden del prevalidador) e informe de validación."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .config import Config

AZUL = PatternFill("solid", fgColor="0C447C")
COLORES_NIVEL = {"ERROR": "F8D7DA", "ALERTA": "FFF3CD", "INFO": "E2E3E5"}


def _formatear(ruta: Path, hojas_texto: dict[str, set[str]] | None = None) -> None:
    wb = load_workbook(ruta)
    for ws in wb.worksheets:
        for c in ws[1]:
            c.font, c.fill = Font(bold=True, color="FFFFFF"), AZUL
            c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.freeze_panes = "A2"
        for i, col in enumerate(ws.columns, 1):
            ancho = max(len(str(c.value or "")) for c in list(col)[:500])
            ws.column_dimensions[get_column_letter(i)].width = min(max(10, ancho + 2), 60)
            if isinstance(col[1].value if len(col) > 1 else None, (int, float)):
                for c in list(col)[1:]:
                    c.number_format = "#,##0"
        if ws.title == "Hallazgos":
            for fila in ws.iter_rows(min_row=2):
                color = COLORES_NIVEL.get(fila[0].value)
                if color:
                    for c in fila:
                        c.fill = PatternFill("solid", fgColor=color)
    wb.save(ruta)


MAX_FILAS_PREVALIDADOR = 5000   # el prevalidador acepta hasta 5.000 registros por archivo


def escribir_formato(fmt: str, df: pd.DataFrame, cfg: Config, salida: Path) -> Path:
    """Un archivo por formato; si supera 5.000 registros se parte en archivos _parte1, _parte2..."""
    spec = cfg.formatos[fmt]
    anio = cfg.parametros["anio_gravable"]
    encabezados = dict(spec["columnas"])
    out = df.rename(columns=encabezados)
    partes = [out.iloc[i:i + MAX_FILAS_PREVALIDADOR] for i in range(0, max(len(out), 1), MAX_FILAS_PREVALIDADOR)]
    ruta = None
    for n, parte in enumerate(partes, 1):
        sufijo = f"_parte{n}" if len(partes) > 1 else ""
        r = salida / f"Formato_{fmt}_v{spec['version']}_AG{anio}{sufijo}.xlsx"
        with pd.ExcelWriter(r, engine="openpyxl") as w:
            parte.to_excel(w, sheet_name=fmt, index=False)
        _formatear(r)
        ruta = ruta or r
    return ruta


def escribir_informe(salida: Path, cfg: Config, hallazgos: list, cuadres: pd.DataFrame,
                     partidas: pd.DataFrame, terceros: pd.DataFrame, generados: dict) -> Path:
    ruta = salida / f"Informe_validacion_exogena_AG{cfg.parametros['anio_gravable']}.xlsx"
    h = pd.DataFrame(hallazgos, columns=["nivel", "categoria", "nit_o_cuenta", "detalle"])
    orden = {"ERROR": 0, "ALERTA": 1, "INFO": 2}
    h = h.drop_duplicates().sort_values(["nivel", "categoria"], key=lambda s: s.map(orden) if s.name == "nivel" else s)

    resumen = [{"indicador": "Formatos generados", "valor": ", ".join(sorted(generados))}]
    for fmt, df in sorted(generados.items()):
        resumen.append({"indicador": f"Registros formato {fmt}", "valor": len(df)})
    for nivel in ("ERROR", "ALERTA", "INFO"):
        resumen.append({"indicador": f"Hallazgos {nivel}", "valor": int((h["nivel"] == nivel).sum())})
    resumen.append({"indicador": "Terceros depurados", "valor": len(terceros)})
    listo = (h["nivel"] == "ERROR").sum() == 0 and not cfg.sin_verificar_usados
    resumen.append({"indicador": "¿Listo para prevalidador?",
                    "valor": "SÍ" if listo else "NO — corregir ERRORES y verificar parámetros"})

    sin_verif = sorted(cfg.sin_verificar_usados)
    with pd.ExcelWriter(ruta, engine="openpyxl") as w:
        pd.DataFrame(resumen).to_excel(w, sheet_name="Resumen", index=False)
        h.to_excel(w, sheet_name="Hallazgos", index=False)
        cuadres.to_excel(w, sheet_name="Cuadres", index=False)
        pd.DataFrame({"parametro_sin_verificar": sin_verif}).to_excel(w, sheet_name="Parametros sin verificar",
                                                                    index=False)
        partidas.to_excel(w, sheet_name="Trazabilidad", index=False)
        terceros.to_excel(w, sheet_name="Terceros", index=False)
    _formatear(ruta)
    return ruta
