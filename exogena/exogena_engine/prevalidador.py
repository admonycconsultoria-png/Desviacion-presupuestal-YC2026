"""Prevalidador de la DIAN (Excel .xlsm) como fuente de verdad del ORDEN DE COLUMNAS.

Cada hoja F#### del prevalidador trae, en filas fijas, la definición de cada columna:
    fila 2 encabezado · fila 3 tipo (N/A) · fila 4 longitud · fila 5 obligatorio (S/N)
    fila 6 tabla de valores válidos · fila 10 etiqueta XML
La hoja DefinicionFormatos trae la versión ("1001(V-10) - ...") y la hoja Tablas los catálogos
(conceptos por formato, países, tipos de documento).

    python -m exogena_engine.prevalidador Prevalidador_AG2025.xlsm            -> informe
    python -m exogena_engine.prevalidador Prevalidador_AG2025.xlsm --guardar  -> config/prevalidadores/<archivo>.json
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

from .utils import clave, texto_dian

FORMATOS = ("1001", "1003", "1005", "1006", "1007", "1008", "1009", "1010", "1011", "1012", "2276")


def _txt(v) -> str:
    """Celda -> texto de una línea (openpyxl deja los saltos de línea como _x000D_)."""
    return "" if v is None else re.sub(r"(?:_x000D_|\s)+", " ", str(v)).strip()


def leer(ruta: str | Path) -> dict:
    """Layouts y catálogos del prevalidador para los formatos que maneja el aplicativo."""
    import openpyxl

    wb = openpyxl.load_workbook(ruta, read_only=True, data_only=True)
    versiones = {}
    for fila in wb["DefinicionFormatos"].iter_rows(min_row=3, values_only=True):
        m = re.match(r"\s*(\d{4})\s*\(V-(\d+)\)", _txt(fila[1]) if len(fila) > 1 else "")
        if m:
            versiones[m.group(1)] = int(m.group(2))

    tablas: dict[str, list[list[str]]] = {}
    filas = list(wb["Tablas"].iter_rows(values_only=True))
    i = 1
    while i < len(filas):
        nombre, n = _txt(filas[i][0]) if filas[i] else "", _txt(filas[i][1]) if filas[i] and len(filas[i]) > 1 else ""
        if nombre.startswith("D") and n.isdigit():
            tablas[nombre] = [[_txt(f[0]), _txt(f[1])] for f in filas[i + 1:i + 1 + int(n)]]
            i += int(n) + 1
        else:
            i += 1

    formatos = {}
    for fmt in FORMATOS:
        hoja = f"F{fmt}"
        if hoja not in wb.sheetnames or fmt not in versiones:
            continue
        r = list(wb[hoja].iter_rows(min_row=1, max_row=10, values_only=True))
        cols, j = [], 1
        while j < len(r[1]) and _txt(r[1][j]) and not _txt(r[1][j]).startswith("|"):
            cols.append({"encabezado": _txt(r[1][j]), "tipo": _txt(r[2][j]),
                         "longitud": int(float(_txt(r[3][j]) or 0)), "obligatorio": _txt(r[4][j]).upper() == "S",
                         "tabla": _txt(r[5][j]), "xml": _txt(r[9][j])})
            j += 1
        conceptos = []
        if cols and cols[0]["xml"] == "cpt" and cols[0]["tabla"] in tablas:
            conceptos = tablas[cols[0]["tabla"]]
        formatos[fmt] = {"version": versiones[fmt], "columnas": cols, "conceptos": conceptos}
    return {
        "nombre": Path(ruta).stem.split("-", 1)[-1] if re.match(r"^[0-9a-f]{8}-", Path(ruta).name) else Path(ruta).stem,
        "formatos": formatos,
        "paises": [[c, texto_dian(n)] for c, n in tablas.get("D_2_107", [])],
        "tipos_documento": tablas.get("DTiposDocto", []),
        "entidad_informante_2276": tablas.get("DEntidadInfor", []),
    }


def estado_columnas(fmt: str, spec: dict, prevalidadores: list[dict]) -> dict:
    """Compara el layout configurado (versión ya resuelta para el año) con el de un prevalidador de la
    MISMA versión. Las columnas deben coincidir en orden; el prevalidador puede traer columnas
    adicionales al final solo si son opcionales (se informan como omitidas)."""
    nuestras = [e for _, e in spec["columnas"]]
    for p in prevalidadores:
        lay = (p.get("formatos") or {}).get(fmt)
        if not lay or int(lay["version"]) != int(spec["version"]):
            continue
        suyas = lay["columnas"]
        difs = []
        for k in range(max(len(nuestras), len(suyas))):
            a = nuestras[k] if k < len(nuestras) else None
            b = suyas[k]["encabezado"] if k < len(suyas) else None
            if a is None and not suyas[k]["obligatorio"]:
                continue
            if a is None or b is None or clave(a) != clave(b):
                difs.append(f"col {k + 1}: configurada '{a or '—'}' / prevalidador '{b or '—'}'")
        omitidas = [c["encabezado"] for c in suyas[len(nuestras):]]
        return {"verificado": not difs, "fuente": p["nombre"], "diferencias": difs, "omitidas": omitidas}
    return {"verificado": False, "fuente": "", "diferencias": [], "omitidas": []}


def main() -> None:
    from . import config as config_mod

    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("xlsm")
    ap.add_argument("--guardar", action="store_true", help="guarda el layout en config/prevalidadores/")
    a = ap.parse_args()
    p = leer(a.xlsm)
    cfg = config_mod.cargar()
    vers = ", ".join(f"{k} v{v['version']}" for k, v in p["formatos"].items())
    print(f"Prevalidador {p['nombre']}: {vers}")
    for fmt, spec in cfg.formatos.items():
        e = estado_columnas(fmt, spec, [p])
        lay = p["formatos"].get(fmt)
        if not lay:
            print(f"  {fmt}: no está en el prevalidador")
        elif int(lay["version"]) != int(spec["version"]):
            print(f"  {fmt}: configurado v{spec['version']} (AG {cfg.parametros['anio_gravable']}), prevalidador v{lay['version']}: no comparable")
        else:
            print(f"  {fmt} v{spec['version']}: {'OK' if e['verificado'] else 'DIFERENTE'}"
                  + (f" (omite opcionales: {'; '.join(e['omitidas'])})" if e["omitidas"] else ""))
            for d in e["diferencias"]:
                print("      " + d)
        if lay and lay["conceptos"]:
            nuestros = set(cfg.conceptos.loc[cfg.conceptos["formato"] == fmt, "concepto"])
            faltan = [c for c, _ in lay["conceptos"] if c not in nuestros]
            if faltan:
                print(f"      conceptos del prevalidador que no están en conceptos.csv: {', '.join(faltan)}")
    if a.guardar:
        destino = config_mod.CONFIG_DIR / "prevalidadores" / f"{p['nombre']}.json"
        destino.parent.mkdir(exist_ok=True)
        destino.write_text(json.dumps(p, ensure_ascii=False, indent=1), encoding="utf-8")
        print(f"Guardado {destino}")


if __name__ == "__main__":
    main()
